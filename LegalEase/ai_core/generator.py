import io
import re
from pathlib import Path
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from fpdf import FPDF
from config import LOGO_PATH, INVERSE_LOGO_PATH

def ensure_default_logos():
    """Generates clean placeholder logo images if missing in the Image/ directory."""
    LOGO_PATH.parent.mkdir(parents=True, exist_ok=True)
    if not LOGO_PATH.exists():
        img = Image.new("RGBA", (300, 80), color=(255, 255, 255, 0))
        draw = ImageDraw.Draw(img)
        draw.rectangle([(10, 10), (290, 70)], outline=(30, 41, 59), width=3)
        draw.text((35, 28), "⚖  LegalEase", fill=(30, 41, 59))
        img.save(LOGO_PATH)
    if not INVERSE_LOGO_PATH.exists():
        img_inv = Image.new("RGBA", (300, 80), color=(255, 255, 255, 0))
        draw_inv = ImageDraw.Draw(img_inv)
        draw_inv.rectangle([(10, 10), (290, 70)], outline=(241, 245, 249), width=3)
        draw_inv.text((35, 28), "⚖  LegalEase", fill=(241, 245, 249))
        img_inv.save(INVERSE_LOGO_PATH)

ensure_default_logos()

def sanitize_text(text: str) -> str:
    """Normalizes Unicode characters, smart quotes, dashes, and bullet points to prevent encoding failures."""
    if not text:
        return ""
    replacements = {
        "“": '"', "”": '"', "‘": "'", "’": "'", "—": "-", "–": "-",
        "…": "...", "•": "* ", "\r\n": "\n", "\r": "\n", "™": "TM",
        "©": "(c)", "®": "(R)"
    }
    for orig, rep in replacements.items():
        text = text.replace(orig, rep)
    return text.strip()

def format_html_preview(text: str) -> str:
    """Converts generated legal markdown into a styled dark-mode HTML preview."""
    sanitized = sanitize_text(text)
    lines = sanitized.split("\n")
    html_parts = []

    for line in lines:
        line = line.strip()
        if not line:
            continue
        if line in ("---", "***", "___"):
            html_parts.append("<hr style='border: 0; border-top: 1px solid #334155; margin: 14px 0;'>")
        elif line.startswith("### "):
            title = line.replace("### ", "").strip()
            html_parts.append(
                f"<h3 style='color: #93c5fd; font-size: 16px; margin-top: 14px; margin-bottom: 6px;'>{title}</h3>"
            )
        elif line.startswith("## "):
            title = line.replace("## ", "").strip()
            html_parts.append(
                f"<h2 style='color: #60a5fa; border-bottom: 1px solid #374151; "
                f"padding-bottom: 6px; margin-top: 14px;'>{title}</h2>"
            )
        elif line.startswith("# "):
            title = line.replace("# ", "").strip()
            html_parts.append(
                f"<h1 style='color: #93c5fd; border-bottom: 2px solid #4b5563; "
                f"padding-bottom: 8px;'>{title}</h1>"
            )
        else:
            formatted_line = re.sub(r"\*\*(.*?)\*\*", r"<strong style='color:#f3f4f6;'>\1</strong>", line)
            html_parts.append(f"<p style='margin-bottom: 8px; line-height: 1.6;'>{formatted_line}</p>")

    return "\n".join(html_parts)

def format_docx(text: str, doc_type: str, terms_input: str = "") -> bytes:
    """Generates an executive-formatted .docx document containing headers, logos, terms tables, and footers."""
    doc = Document()
    sanitized = sanitize_text(text)

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        # Footer
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f_run = f_p.add_run("LegalEase Inc | contact@legalease.com | All Rights Reserved.")
        f_run.font.name = "Times New Roman"
        f_run.font.size = Pt(9)
        f_run.font.color.rgb = RGBColor(128, 128, 128)

    # Document Header Logo
    if LOGO_PATH.exists():
        logo_para = doc.add_paragraph()
        logo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo_para.add_run().add_picture(str(LOGO_PATH), width=Inches(2.2))

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(doc_type.upper())
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(16)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)

    # Render Body
    paragraphs = sanitized.split("\n\n")
    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        if para.startswith("## ") or para.startswith("# "):
            clean_head = para.lstrip("#").strip()
            if clean_head.lower() == doc_type.lower():
                continue
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            h_run = h.add_run(clean_head)
            h_run.font.name = "Times New Roman"
            h_run.font.size = Pt(12)
            h_run.bold = True
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(6)
            tokens = re.split(r"(\*\*.*?\*\*)", para)
            for token in tokens:
                if token.startswith("**") and token.endswith("**"):
                    r = p.add_run(token[2:-2])
                    r.bold = True
                else:
                    r = p.add_run(token)
                r.font.name = "Times New Roman"
                r.font.size = Pt(11)

    # Append Terms & Conditions Table if input is present
    if terms_input:
        term_items = [t.strip() for t in terms_input.split(";") if t.strip()]
        if term_items:
            t_head = doc.add_paragraph()
            t_head.paragraph_format.space_before = Pt(14)
            t_head_run = t_head.add_run("Summary of Specific Clauses & Terms")
            t_head_run.font.name = "Times New Roman"
            t_head_run.font.size = Pt(12)
            t_head_run.bold = True

            table = doc.add_table(rows=1, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False

            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = "Item #"
            hdr_cells[1].text = "Agreed Stipulation / Condition"
            hdr_cells[0].paragraphs[0].runs[0].font.bold = True
            hdr_cells[1].paragraphs[0].runs[0].font.bold = True

            for idx, item in enumerate(term_items, start=1):
                row_cells = table.add_row().cells
                row_cells[0].text = f"Clause {idx}"
                row_cells[1].text = item
                row_cells[0].paragraphs[0].runs[0].font.name = "Times New Roman"
                row_cells[1].paragraphs[0].runs[0].font.name = "Times New Roman"

    # Export docx buffer
    target_stream = io.BytesIO()
    doc.save(target_stream)
    return target_stream.getvalue()

class LegalEasePDF(FPDF):
    def __init__(self, doc_type: str = "Legal Document"):
        super().__init__()
        self.doc_type = doc_type

    def header(self):
        if LOGO_PATH.exists():
            self.image(str(LOGO_PATH), x=85, y=10, w=40)
            self.ln(18)
        self.set_font("Helvetica", "B", 13)
        self.cell(0, 8, self.doc_type.upper(), align="C", new_x="LMARGIN", new_y="NEXT")
        self.set_draw_color(200, 200, 200)
        self.line(10, self.get_y(), 200, self.get_y())
        self.ln(5)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(128, 128, 128)
        self.cell(0, 10, f"LegalEase Inc | contact@legalease.com | All Rights Reserved.   (Page {self.page_no()})", align="C")

def format_pdf(text: str, doc_type: str) -> bytes:
    """Generates a clean PDF document with embedded headers, logos, and footers on every page."""
    pdf = LegalEasePDF(doc_type=doc_type)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()
    pdf.set_margins(18, 18, 18)

    sanitized = sanitize_text(text)
    paragraphs = sanitized.split("\n\n")

    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        if para.startswith("## ") or para.startswith("# "):
            heading_text = para.lstrip("#").strip()
            if heading_text.lower() == doc_type.lower():
                continue
            pdf.ln(3)
            pdf.set_font("Helvetica", "B", 11)
            pdf.set_text_color(20, 30, 50)
            pdf.multi_cell(0, 6, heading_text.encode("latin-1", "replace").decode("latin-1"))
            pdf.ln(1)
        else:
            pdf.set_font("Helvetica", "", 10)
            pdf.set_text_color(40, 40, 40)
            clean_p = re.sub(r"\*\*(.*?)\*\*", r"\1", para)
            pdf.multi_cell(0, 5.5, clean_p.encode("latin-1", "replace").decode("latin-1"))
            pdf.ln(2.5)

    return bytes(pdf.output())